
import csv
import json
import io
import os
import re
import sqlite3
import tempfile
import threading
import time
import uuid
from datetime import date, datetime
from pathlib import Path

from flask import jsonify, request
from werkzeug.utils import secure_filename

MAX_UPLOAD_BYTES = 15 * 1024 * 1024
MAX_ROWS = 20000
MAX_COLUMNS = 100
MAX_RESULT_ROWS = 200
WORKSPACE_TTL_SECONDS = 60 * 60
_WORKSPACE_ROOT = tempfile.mkdtemp(prefix="sqlxpert_imports_")
_WORKSPACES = {}
_LOCK = threading.Lock()
_ALLOWED_EXTENSIONS = {".csv", ".json", ".xlsx", ".xlsm"}


def _safe_identifier(value, fallback="table"):
    value = re.sub(r"[^A-Za-z0-9_]", "_", str(value or "").strip())
    if not value or not re.match(r"^[A-Za-z_]", value):
        value = "data_" + value
    value = value[:60] or fallback
    if value.lower().startswith("sqlite_"):
        value = "data_" + value
    return value


def _unique_names(names):
    used, result = set(), []
    for index, raw in enumerate(names, 1):
        name = _safe_identifier(raw, f"column_{index}")
        candidate, suffix = name, 2
        while candidate.lower() in used:
            candidate = f"{name}_{suffix}"
            suffix += 1
        used.add(candidate.lower())
        result.append(candidate)
    return result


def _clean_value(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (str, int, float)):
        if isinstance(value, float) and (value != value or value in (float("inf"), float("-inf"))):
            return None
        return value
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def _infer_type(values):
    present = [v for v in values if v is not None and str(v).strip() != ""]
    if not present:
        return "TEXT"
    def as_int(v):
        if isinstance(v, bool): return True
        if isinstance(v, int): return True
        if isinstance(v, float): return v.is_integer()
        return bool(re.fullmatch(r"[+-]?\d+", str(v).strip()))
    def as_float(v):
        try:
            float(str(v).strip())
            return True
        except (ValueError, TypeError):
            return False
    if all(as_int(v) for v in present): return "INTEGER"
    if all(as_float(v) for v in present): return "REAL"
    return "TEXT"


def _rows_from_csv(raw):
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Could not decode this CSV file. Save it as UTF-8 and try again.")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    records = list(reader)
    if not records:
        raise ValueError("The file is empty.")
    headers = [str(x).strip() for x in records[0]]
    if not any(headers):
        raise ValueError("The first CSV row must contain column headers.")
    names = _unique_names([h or f"column_{i+1}" for i, h in enumerate(headers)])
    rows = []
    for row in records[1:]:
        if not row or not any(str(v).strip() for v in row):
            continue
        row = row[:len(names)] + [None] * max(0, len(names) - len(row))
        rows.append([_clean_value(v if v != "" else None) for v in row])
        if len(rows) > MAX_ROWS:
            raise ValueError(f"Files are limited to {MAX_ROWS:,} data rows.")
    return names, rows


def _rows_from_json(raw):
    try:
        data = json.loads(raw.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"Invalid JSON file: {exc}")
    if isinstance(data, dict):
        # Common API format: {"data": [{...}, ...]}.
        for key in ("data", "rows", "records", "items"):
            if isinstance(data.get(key), list):
                data = data[key]
                break
        else:
            data = [data]
    if not isinstance(data, list) or not data:
        raise ValueError("JSON must contain an object or a non-empty array of objects.")
    if not all(isinstance(row, dict) for row in data):
        raise ValueError("JSON arrays must contain objects with named fields, for example [{\"name\":\"Asha\",\"marks\":92}].")
    original_keys = []
    for row in data:
        for key in row.keys():
            if str(key) not in original_keys:
                original_keys.append(str(key))
    if not original_keys:
        raise ValueError("No columns were found in the JSON objects.")
    names = _unique_names(original_keys)
    rows = [[_clean_value(row.get(key)) for key in original_keys] for row in data[:MAX_ROWS + 1]]
    if len(data) > MAX_ROWS:
        raise ValueError(f"Files are limited to {MAX_ROWS:,} data rows.")
    return names, rows


def _rows_from_excel(raw):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("Excel support requires openpyxl. In your project terminal run: python -m pip install openpyxl")
    import io
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read this Excel workbook: {exc}")
    try:
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header or not any(v is not None and str(v).strip() for v in header):
            raise ValueError("The first worksheet must have a header row.")
        names = _unique_names([str(v).strip() if v is not None and str(v).strip() else f"column_{i+1}" for i, v in enumerate(header)])
        rows = []
        for row in iterator:
            vals = list(row[:len(names)]) + [None] * max(0, len(names) - len(row))
            vals = [_clean_value(v) for v in vals]
            if not any(v is not None and str(v).strip() for v in vals):
                continue
            rows.append(vals)
            if len(rows) > MAX_ROWS:
                raise ValueError(f"Files are limited to {MAX_ROWS:,} data rows.")
        return names, rows
    finally:
        workbook.close()


def _load_records(filename, raw):
    suffix = Path(filename).suffix.lower()
    if suffix not in _ALLOWED_EXTENSIONS:
        raise ValueError("Supported formats are CSV (.csv), JSON (.json), and Excel (.xlsx, .xlsm).")
    if not raw:
        raise ValueError("The uploaded file is empty.")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError("Maximum upload size is 15 MB.")
    if suffix == ".csv":
        columns, rows = _rows_from_csv(raw)
    elif suffix == ".json":
        columns, rows = _rows_from_json(raw)
    else:
        columns, rows = _rows_from_excel(raw)
    if len(columns) > MAX_COLUMNS:
        raise ValueError(f"Files are limited to {MAX_COLUMNS} columns.")
    types = [_infer_type([row[i] for row in rows]) for i in range(len(columns))]
    table = _safe_identifier(Path(secure_filename(filename)).stem or "imported_data")
    return table, columns, types, rows, suffix


def _cleanup_expired():
    cutoff = time.time() - WORKSPACE_TTL_SECONDS
    with _LOCK:
        expired = [key for key, value in _WORKSPACES.items() if value["created"] < cutoff]
        for key in expired:
            info = _WORKSPACES.pop(key, None)
            if info:
                try: os.remove(info["path"])
                except OSError: pass


def register_import_workspace(app):
    """Register the import workspace routes on the existing Flask application."""
    if app.config.get("MAX_CONTENT_LENGTH") is None or app.config.get("MAX_CONTENT_LENGTH") > MAX_UPLOAD_BYTES:
        app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES

    @app.errorhandler(413)
    def imported_file_too_large(_error):
        return jsonify({"error": "Maximum upload size is 15 MB."}), 413

    @app.post("/api/import-workspace/upload")
    def import_workspace_upload():
        _cleanup_expired()
        upload = request.files.get("file")
        if upload is None or not upload.filename:
            return jsonify({"error": "Choose a CSV, JSON, or Excel file first."}), 400
        try:
            raw = upload.read(MAX_UPLOAD_BYTES + 1)
            if len(raw) > MAX_UPLOAD_BYTES:
                raise ValueError("Maximum upload size is 15 MB.")
            table, columns, types, rows, file_type = _load_records(upload.filename, raw)
            workspace_id = uuid.uuid4().hex
            db_path = os.path.join(_WORKSPACE_ROOT, workspace_id + ".sqlite3")
            conn = sqlite3.connect(db_path)
            try:
                definitions = ", ".join('"' + name.replace('"', '""') + '" ' + kind for name, kind in zip(columns, types))
                conn.execute(f'CREATE TABLE "{table}" ({definitions})')
                if rows:
                    placeholders = ",".join("?" for _ in columns)
                    conn.executemany(f'INSERT INTO "{table}" VALUES ({placeholders})', rows)
                conn.commit()
            finally:
                conn.close()
            schema = [{"name": name, "type": kind} for name, kind in zip(columns, types)]
            with _LOCK:
                _WORKSPACES[workspace_id] = {
                    "path": db_path, "created": time.time(), "tables": {table: schema},
                    "row_count": len(rows), "filename": secure_filename(upload.filename), "file_type": file_type
                }
            return jsonify({"success": True, "workspace_id": workspace_id, "filename": secure_filename(upload.filename),
                            "table": table, "tables": [{"name": table, "columns": schema}],
                            "row_count": len(rows), "column_count": len(columns),
                            "message": f"Imported {len(rows):,} rows into temporary table '{table}'."})
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        except Exception as exc:
            app.logger.exception("Import workspace failed")
            return jsonify({"error": f"Could not import this file: {exc}"}), 400

    @app.post("/api/import-workspace/execute")
    def import_workspace_execute():
        data = request.get_json(silent=True) or {}
        workspace_id = str(data.get("workspace_id", ""))
        sql = str(data.get("sql", "")).strip()
        with _LOCK:
            workspace = _WORKSPACES.get(workspace_id)
        if not workspace:
            return jsonify({"error": "This temporary data workspace expired or was cleared. Upload the file again."}), 404
        if not sql:
            return jsonify({"error": "Enter a SELECT query first."}), 400
        # Single read-only SELECT only. SQLite authorizer provides a second guard.
        cleaned = re.sub(r"--[^\n]*|/\*[\s\S]*?\*/", "", sql).strip()
        without_trailing = cleaned[:-1] if cleaned.endswith(";") else cleaned
        if not re.match(r"(?is)^SELECT\b", cleaned) or ";" in without_trailing:
            return jsonify({"error": "Imported-data mode supports one read-only SELECT query at a time."}), 400
        start = time.perf_counter()
        conn = sqlite3.connect(workspace["path"])
        conn.row_factory = sqlite3.Row
        allowed_tables = {name.lower() for name in workspace["tables"]}
        def authorizer(action, arg1, arg2, _database, _source):
            if action == sqlite3.SQLITE_SELECT:
                return sqlite3.SQLITE_OK
            if action == sqlite3.SQLITE_READ:
                return sqlite3.SQLITE_OK if str(arg1 or "").lower() in allowed_tables else sqlite3.SQLITE_DENY
            if action in (sqlite3.SQLITE_FUNCTION,):
                return sqlite3.SQLITE_OK
            return sqlite3.SQLITE_DENY
        try:
            conn.set_authorizer(authorizer)
            started = time.perf_counter()
            conn.set_progress_handler(lambda: 1 if time.perf_counter() - started > 2.0 else 0, 10000)
            cursor = conn.execute(sql)
            raw_rows = cursor.fetchmany(MAX_RESULT_ROWS + 1)
            truncated = len(raw_rows) > MAX_RESULT_ROWS
            raw_rows = raw_rows[:MAX_RESULT_ROWS]
            columns = [d[0] for d in (cursor.description or [])]
            rows = [dict(row) for row in raw_rows]
            duration = (time.perf_counter() - start) * 1000
            with _LOCK:
                if workspace_id in _WORKSPACES:
                    _WORKSPACES[workspace_id]["created"] = time.time()
            return jsonify({"success": True, "command": "SELECT", "columns": columns, "rows": rows,
                            "row_count": len(rows), "truncated": truncated, "duration_ms": round(duration, 3),
                            "message": "Query executed against your temporary imported dataset."})
        except sqlite3.Error as exc:
            message = "Query timed out. Try a smaller query." if "interrupt" in str(exc).lower() else str(exc)
            return jsonify({"error": message}), 400
        finally:
            conn.close()

    @app.post("/api/import-workspace/explain")
    def import_workspace_explain():
        data = request.get_json(silent=True) or {}
        workspace_id = str(data.get("workspace_id", ""))
        sql = str(data.get("sql", "")).strip()
        with _LOCK:
            workspace = _WORKSPACES.get(workspace_id)
        if not workspace:
            return jsonify({"error": "This temporary data workspace expired or was cleared. Upload the file again."}), 404
        cleaned = re.sub(r"--[^\n]*|/\*[\s\S]*?\*/", "", sql).strip()
        without_trailing = cleaned[:-1] if cleaned.endswith(";") else cleaned
        if not re.match(r"(?is)^SELECT\b", cleaned) or ";" in without_trailing:
            return jsonify({"error": "Explain Plan supports one read-only SELECT query at a time."}), 400
        allowed_tables = {name.lower() for name in workspace["tables"]}
        def authorizer(action, arg1, arg2, _database, _source):
            if action == sqlite3.SQLITE_SELECT:
                return sqlite3.SQLITE_OK
            if action == sqlite3.SQLITE_READ:
                return sqlite3.SQLITE_OK if str(arg1 or "").lower() in allowed_tables else sqlite3.SQLITE_DENY
            if action == sqlite3.SQLITE_FUNCTION:
                return sqlite3.SQLITE_OK
            return sqlite3.SQLITE_DENY
        conn = sqlite3.connect(workspace["path"])
        try:
            conn.set_authorizer(authorizer)
            plan_rows = conn.execute("EXPLAIN QUERY PLAN " + sql).fetchall()
            plan = []
            for row in plan_rows:
                detail = row[3]
                if "SCAN" in detail.upper():
                    explanation = "SQLite scans rows from an imported table or subquery."
                elif "SEARCH" in detail.upper():
                    explanation = "SQLite uses a search access path for this plan step."
                else:
                    explanation = "An operation in SQLite's physical query plan."
                plan.append({"detail": detail, "explanation": explanation})
            with _LOCK:
                if workspace_id in _WORKSPACES:
                    _WORKSPACES[workspace_id]["created"] = time.time()
            return jsonify({"plan": plan, "note": "This is SQLite's actual plan for the temporary imported dataset. It is not a cost estimate."})
        except sqlite3.Error as exc:
            return jsonify({"error": str(exc)}), 400
        finally:
            conn.close()

    @app.post("/api/import-workspace/close")
    def import_workspace_close():
        data = request.get_json(silent=True) or {}
        workspace_id = str(data.get("workspace_id", ""))
        with _LOCK:
            workspace = _WORKSPACES.pop(workspace_id, None)
        if workspace:
            try: os.remove(workspace["path"])
            except OSError: pass
        return jsonify({"success": True, "message": "Temporary imported dataset cleared."})
